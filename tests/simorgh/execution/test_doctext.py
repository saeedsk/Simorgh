"""Files that are not plain text (execution/doctext.py).

`read_file` understood source and PDF; everything else came back as
binary noise, so a spreadsheet of the very data a task was about was
unreachable. Every dependency here is optional, so each test that needs
one skips without it -- and the "not installed" path is asserted too,
because a missing library must produce an install hint, never an
exception.
"""

from __future__ import annotations

import importlib.util
import io
import tempfile
import unittest
import unittest.mock
from pathlib import Path

from simorgh.execution import doctext
from simorgh.execution.config import Config
from simorgh.execution.doctext import (
    csv_to_text, document_to_text, docx_to_text, image_to_text,
    looks_like_docx, looks_like_image, looks_like_xlsx, xlsx_to_text,
)

_HAS_DOCX = importlib.util.find_spec("docx") is not None
_HAS_XLSX = importlib.util.find_spec("openpyxl") is not None
_HAS_PIL = importlib.util.find_spec("PIL") is not None


def _xlsx_bytes(rows, title="Sheet1"):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = title
    for row in rows:
        ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _docx_bytes(paragraphs, table=None):
    import docx

    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _png_bytes(size=(120, 40)):
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", size, "white").save(buffer, format="PNG")
    return buffer.getvalue()


class SniffingTestCase(unittest.TestCase):
    @unittest.skipUnless(_HAS_XLSX and _HAS_DOCX, "openpyxl/python-docx not installed")
    def test_docx_and_xlsx_are_told_apart_despite_both_being_zips(self):
        docx_data, xlsx_data = _docx_bytes(["hi"]), _xlsx_bytes([["a"]])
        self.assertTrue(looks_like_docx(docx_data))
        self.assertFalse(looks_like_xlsx(docx_data))
        self.assertTrue(looks_like_xlsx(xlsx_data))
        self.assertFalse(looks_like_docx(xlsx_data))

    @unittest.skipUnless(_HAS_PIL, "Pillow not installed")
    def test_an_image_is_recognised_by_its_magic_bytes(self):
        self.assertTrue(looks_like_image(_png_bytes()))

    def test_plain_text_is_nobody_here_business(self):
        self.assertIsNone(document_to_text(b"def f():\n    pass\n", name="a.py"))
        self.assertFalse(looks_like_docx(b"not a zip"))
        self.assertFalse(looks_like_image(b"hello"))

    def test_a_truncated_zip_does_not_raise(self):
        self.assertFalse(looks_like_docx(b"PK\x03\x04truncated"))


class ContentTestCase(unittest.TestCase):
    @unittest.skipUnless(_HAS_DOCX, "python-docx not installed")
    def test_a_docx_yields_its_paragraphs_and_tables(self):
        data = _docx_bytes(["Almaden valley report"], table=[["address", "price"], ["Via Corta", "4218000"]])
        text, problem = docx_to_text(data, max_chars=20_000)
        self.assertEqual(problem, "")
        self.assertIn("Almaden valley report", text)
        self.assertIn("address\tprice", text)

    @unittest.skipUnless(_HAS_XLSX, "openpyxl not installed")
    def test_an_xlsx_yields_rows_with_the_sheet_named(self):
        data = _xlsx_bytes([["address", "price"], ["Via Corta", 4218000]], title="Listings")
        text, problem = xlsx_to_text(data, max_rows=200, max_chars=20_000)
        self.assertEqual(problem, "")
        self.assertIn("# sheet: Listings", text)
        self.assertIn("Via Corta\t4218000", text)

    @unittest.skipUnless(_HAS_XLSX, "openpyxl not installed")
    def test_the_row_cap_is_honoured_and_says_so(self):
        data = _xlsx_bytes([[i] for i in range(50)])
        text, _ = xlsx_to_text(data, max_rows=5, max_chars=20_000)
        self.assertIn("first 5 rows", text)

    def test_a_csv_becomes_tab_separated_rows(self):
        text, problem = csv_to_text(b"a,b\n1,2\n", max_rows=10, max_chars=1000)
        self.assertEqual(problem, "")
        self.assertEqual(text, "a\tb\n1\t2")

    def test_a_csv_row_cap_is_honoured(self):
        raw = b"h\n" + b"\n".join(str(i).encode() for i in range(50))
        text, _ = csv_to_text(raw, max_rows=5, max_chars=1000)
        self.assertIn("first 5 rows", text)

    @unittest.skipUnless(_HAS_PIL, "Pillow not installed")
    def test_an_image_reports_its_shape_and_does_not_invent_content(self):
        # Claiming to know what is in a photo from its bytes would be
        # exactly the confident fiction this project keeps catching.
        text, problem = image_to_text(_png_bytes((120, 40)), max_chars=2000, ocr=False)
        self.assertEqual(problem, "")
        self.assertIn("120x40", text)
        self.assertIn("PNG", text)

    @unittest.skipUnless(_HAS_PIL, "Pillow not installed")
    def test_missing_ocr_is_stated_not_silently_skipped(self):
        text, _ = image_to_text(_png_bytes(), max_chars=2000, ocr=True)
        # Either OCR ran, or the reason it did not is in the text.
        self.assertTrue("OCR" in text or "ocr" in text)

    def test_a_corrupt_file_is_a_problem_not_an_exception(self):
        if _HAS_DOCX:
            text, problem = docx_to_text(b"PK\x03\x04garbage", max_chars=100)
            self.assertEqual(text, "")
            self.assertTrue(problem)

    def test_a_missing_library_gives_an_install_hint(self):
        with unittest.mock.patch.dict("sys.modules", {"docx": None}):
            text, problem = docx_to_text(b"PK\x03\x04", max_chars=100)
        self.assertEqual(text, "")
        self.assertIn("python-docx", problem)


class ZipBombTestCase(unittest.TestCase):
    """A .docx/.xlsx is a zip, and neither python-docx nor openpyxl
    checks how much a member expands before decompressing and parsing
    it fully into memory. Confirmed live (observer, 2026-09-09): a 4.9
    MB crafted .docx decompressed to 2 GB and was still climbing past
    5.7 GB of RSS after 20 seconds; a 5.4 MB crafted .xlsx decompressed
    to 1.6 GB and had not returned after 120 seconds despite
    `max_rows=200` (openpyxl's read_only iteration does not actually
    stop cheap work at the row cap). Both ran synchronously inside an
    `async def tool.run()` with no await point, so Execution's
    `asyncio.wait_for` timeout could never fire either -- the whole
    event loop was starved for as long as the parse ran. These tests
    build a MUCH smaller bomb (well under 1 MB compressed, over the 50
    MB uncompressed guard) so the suite stays fast, and assert the
    refusal is instant and never invokes the real parser."""

    @staticmethod
    def _zip_with_bomb_member(member_name: str, real_members: dict, bomb_size: int) -> bytes:
        import zipfile

        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
            for name, data in real_members.items():
                if name == member_name:
                    archive.writestr(name, b"A" * bomb_size)
                else:
                    archive.writestr(name, data)
        return buffer.getvalue()

    @unittest.skipUnless(_HAS_DOCX, "python-docx not installed")
    def test_a_docx_decompression_bomb_is_refused_before_parsing(self):
        import zipfile

        real = _docx_bytes(["seed"])
        with zipfile.ZipFile(io.BytesIO(real)) as archive:
            members = {name: archive.read(name) for name in archive.namelist()}
        bomb = self._zip_with_bomb_member("word/document.xml", members, bomb_size=200_000_000)
        self.assertLess(len(bomb), 1_000_000, "the compressed bomb itself should stay tiny")

        text, problem = docx_to_text(bomb, max_chars=2000)
        self.assertEqual(text, "")
        self.assertIn("decompression bomb", problem)

    @unittest.skipUnless(_HAS_XLSX, "openpyxl not installed")
    def test_an_xlsx_decompression_bomb_is_refused_before_parsing(self):
        import zipfile

        real = _xlsx_bytes([["a"]])
        with zipfile.ZipFile(io.BytesIO(real)) as archive:
            members = {name: archive.read(name) for name in archive.namelist()}
        bomb = self._zip_with_bomb_member("xl/worksheets/sheet1.xml", members, bomb_size=200_000_000)
        self.assertLess(len(bomb), 1_000_000, "the compressed bomb itself should stay tiny")

        text, problem = xlsx_to_text(bomb, max_rows=200, max_chars=2000)
        self.assertEqual(text, "")
        self.assertIn("decompression bomb", problem)

    @unittest.skipUnless(_HAS_DOCX, "python-docx not installed")
    def test_an_ordinary_docx_under_the_guard_still_parses(self):
        data = _docx_bytes(["ordinary paragraph, nothing to see here"])
        text, problem = docx_to_text(data, max_chars=2000)
        self.assertEqual(problem, "")
        self.assertIn("ordinary paragraph", text)


class ReadFileIntegrationTestCase(unittest.TestCase):
    """Through the real `read_file` path, not the helpers directly."""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.root = Path(self._tmp.name)
        (self.root / "docs").mkdir()
        self.addCleanup(self._tmp.cleanup)

    def _read(self, name):
        from simorgh.execution import pathsafety

        return pathsafety.safe_read_file(
            self.root, f"docs/{name}", readable_roots=Config().readable_roots)

    @unittest.skipUnless(_HAS_XLSX, "openpyxl not installed")
    def test_read_file_opens_a_spreadsheet(self):
        (self.root / "docs" / "d.xlsx").write_bytes(_xlsx_bytes([["a", "b"], [1, 2]]))
        self.assertIn("a\tb", self._read("d.xlsx"))

    @unittest.skipUnless(_HAS_DOCX, "python-docx not installed")
    def test_magic_bytes_win_over_a_misleading_extension(self):
        # A .docx renamed .bin still reads: the sniffer looks at the
        # bytes before it looks at the name.
        (self.root / "docs" / "report.bin").write_bytes(_docx_bytes(["hello there"]))
        self.assertIn("hello there", self._read("report.bin"))

    def test_source_files_are_completely_unaffected(self):
        (self.root / "docs" / "a.py").write_text("def f():\n    return 1\n")
        self.assertEqual(self._read("a.py"), "def f():\n    return 1\n")
